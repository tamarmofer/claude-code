#!/usr/bin/env python3
"""Dump the assumption-bearing parts of a real estate / PE Excel model to stdout.

Goal: surface the inputs a skeptic attacks (cap rates, rent, growth, exit, leverage,
hold, DSCR, opex) without dumping the entire workbook into context.

Usage:
    python3 read_model.py path/to/model.xlsx [--all]

Without --all, prints: sheet inventory, named ranges, and every cell whose label or
value looks return-driving. With --all, also prints a compact grid of every non-empty
cell per sheet (use only when the targeted scan misses something).
"""
import sys
import re

# Keywords that flag a row/label as return-driving. Lowercased substring match.
KEYWORDS = [
    "cap rate", "cap ", "exit cap", "going-in", "going in", "terminal",
    "rent", "psf", "$/sf", "per sf", "per unit", "per key", "revpar", "adr",
    "growth", "escalation", "bump", "mark-to-market", "loss to lease", "ltl",
    "irr", "moic", "multiple", "equity multiple", "return", "hurdle", "promote",
    "noi", "egi", "opex", "operating expense", "expense ratio", "tax", "insurance",
    "vacancy", "occupancy", "absorption", "lease-up", "lease up", "downtime",
    "ti", "tenant improvement", "lc", "leasing commission", "capex", "reserve",
    "ltv", "ltc", "dscr", "debt yield", "loan", "refi", "refinance", "coupon",
    "rate", "amort", "interest", "proceeds", "basis", "price", "purchase",
    "replacement cost", "hold", "exit year", "reversion", "residual", "sale",
    "contingency", "gmp", "stabiliz", "wault", "walt",
]


def looks_relevant(text):
    if text is None:
        return False
    t = str(text).strip().lower()
    if not t:
        return False
    return any(k in t for k in KEYWORDS)


def fmt(v):
    if isinstance(v, float):
        # keep it readable; don't force decimals on integers
        return f"{v:,.4g}"
    return str(v)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dump_all = "--all" in sys.argv
    if not args:
        print("usage: read_model.py path/to/model.xlsx [--all]", file=sys.stderr)
        sys.exit(2)
    path = args[0]

    try:
        import openpyxl
    except ImportError:
        print("openpyxl not installed. Install with:\n  pip install openpyxl\n"
              "Or unzip the .xlsx and read xl/worksheets/*.xml + xl/sharedStrings.xml.",
              file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    print(f"# MODEL: {path}")
    print(f"# Sheets: {', '.join(wb.sheetnames)}\n")

    # Named ranges — often the cleanest list of key assumptions.
    try:
        names = list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") \
            else [n.name for n in wb.defined_names.definedName]
        if names:
            print("## Named ranges")
            for n in sorted(names):
                print(f"  - {n}")
            print()
    except Exception:
        pass

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=False))
        if not rows:
            continue
        flagged = []
        for row in rows:
            cells = [c for c in row if c.value is not None]
            if not cells:
                continue
            # A row is interesting if any cell's text matches a keyword.
            label = " | ".join(fmt(c.value) for c in cells)
            if any(looks_relevant(c.value) for c in cells):
                coord = cells[0].coordinate
                flagged.append(f"  [{coord}] {label}")
        if flagged:
            print(f"## Sheet: {ws.title} — return-driving rows")
            print("\n".join(flagged[:200]))
            if len(flagged) > 200:
                print(f"  ... ({len(flagged)-200} more flagged rows)")
            print()

    if dump_all:
        print("\n# ---- FULL NON-EMPTY GRID (--all) ----")
        for ws in wb.worksheets:
            print(f"\n## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=False):
                cells = [(c.coordinate, c.value) for c in row if c.value is not None]
                if cells:
                    print("  " + " | ".join(f"{co}={fmt(v)}" for co, v in cells))


if __name__ == "__main__":
    main()
